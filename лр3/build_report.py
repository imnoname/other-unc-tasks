"""Build the editable GUAP report for laboratory work 3."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(r"C:\Users\theju\OneDrive\Документы\ChatGPT\учеба\_guap_refs\lab_title.docx")
OUTPUT = ROOT / "report.docx"
RESULTS = ROOT / "results"


def set_run_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Обновите поле содержания в Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(0)
    for name, size in (("Heading 1", 14), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(1.25)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(12)
        style.paragraph_format.keep_with_next = True
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph(paragraph)
    add_field(paragraph, "PAGE")
    for run in paragraph.runs:
        set_run_font(run, size=12)


def add_title_page(doc):
    doc.paragraphs[1].text = "КАФЕДРА № 43"
    doc.paragraphs[7].text = "Санкт-Петербург 2026"
    for idx in (2, 3):
        paragraph = doc.paragraphs[idx]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
    for idx in (0, 1, 7):
        for run in doc.paragraphs[idx].runs:
            set_run_font(run, size=14, bold=(idx == 0))
    teacher_table = doc.tables[0]
    set_cell_text(teacher_table.cell(0, 4), "Скобцов Ю.А.", size=12)
    set_cell_text(teacher_table.cell(1, 4), "инициалы, фамилия", size=10)
    title_table = doc.tables[1]
    set_cell_text(title_table.cell(0, 0), "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3", size=13, bold=True)
    set_cell_text(title_table.cell(1, 0), "Решение задач комбинаторной оптимизации с помощью генетических алгоритмов\nна примере задачи укладки рюкзака\nВариант 4", size=10)
    set_cell_text(title_table.cell(2, 0), "по курсу: Эволюционные методы проектирования программно-информационных систем", size=10)
    for row in list(title_table.rows)[3:]:
        title_table._tbl.remove(row._tr)
    student_table = doc.tables[2]
    for paragraph in student_table.cell(0, 0).paragraphs:
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=12)
    set_cell_text(student_table.cell(0, 1), "4332К", size=12)
    set_cell_text(student_table.cell(0, 5), "Могилатов С.И.", size=12)
    set_cell_text(student_table.cell(1, 3), "подпись, дата", size=10)
    set_cell_text(student_table.cell(1, 5), "инициалы, фамилия", size=10)
    doc.add_page_break()


def add_heading(doc, text, level=1, page_break=True):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=14, bold=True)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    set_run_font(run, size=14)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=12)


def add_image(doc, filename, caption, width_cm=15.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.add_run().add_picture(str(RESULTS / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")


def add_table(doc, caption, headers, rows, font_size=10):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(caption)
    set_run_font(run, size=12, bold=True)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], str(header), size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=font_size)
    doc.add_paragraph()
    return table


def add_code_listing(doc, filename, title):
    add_heading(doc, title, level=1, page_break=True)
    source = (ROOT / filename).read_text(encoding="utf-8")
    for line in source.splitlines():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=8.5)


def load_results():
    summary = json.loads((RESULTS / "results_summary.json").read_text(encoding="utf-8"))
    with (RESULTS / "baseline.csv").open(encoding="utf-8") as stream:
        baseline = list(csv.DictReader(stream))
    with (RESULTS / "parameter_study.csv").open(encoding="utf-8") as stream:
        study = list(csv.DictReader(stream))
    return summary, baseline, study


def mean(rows, key):
    return sum(float(row[key]) for row in rows) / len(rows)


def build():
    doc = Document(str(TEMPLATE))
    configure_document(doc)
    add_title_page(doc)
    contents_heading = doc.add_paragraph()
    contents_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contents_heading.paragraph_format.first_line_indent = Cm(0)
    contents_heading.paragraph_format.space_after = Pt(12)
    contents_run = contents_heading.add_run("СОДЕРЖАНИЕ")
    set_run_font(contents_run, size=14, bold=True)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_toc(toc)

    summary, baseline, study = load_results()
    p04 = summary["problems"]["P04"]
    set7 = summary["problems"]["Набор 7"]

    add_heading(doc, "ВВЕДЕНИЕ", level=1)
    add_body(doc, "Генетические алгоритмы позволяют искать решения в больших дискретных пространствах, используя имитацию естественного отбора. В лабораторной работе исследуется задача 0/1-рюкзака, в которой необходимо выбрать набор предметов максимальной стоимости при ограниченной вместимости.")
    add_body(doc, "Цель работы - реализовать генетический алгоритм с двоичным кодированием для задач укладки рюкзака, сравнить результат с точным оптимумом, исследовать влияние основных параметров и проанализировать алгоритм восстановления недопустимых решений.")
    add_body(doc, "В работе рассмотрены тест P04 простой сложности и набор 7 повышенной сложности для варианта 4. Приведены постановка задачи, теоретические сведения, описание программы, результаты экспериментов, графики сходимости и ответ на контрольный вопрос.")

    add_heading(doc, "1 ЗАДАНИЕ И ПОСТАНОВКА ЗАДАЧИ", level=1)
    add_heading(doc, "1.1 Индивидуальное задание", level=2, page_break=False)
    add_body(doc, "Согласно методическим указаниям для варианта 4 требуется использовать двоичное кодирование решения. Для задания простой сложности выбирается тестовый набор 4 (P04), для задания повышенной сложности - набор 7. Необходимо реализовать генетический алгоритм, сравнить решение с оптимальным, построить графики сходимости при разных параметрах, проанализировать время работы и точность, а также ответить на контрольный вопрос 4.")
    add_table(doc, "Таблица 1 - Параметры индивидуального задания", ["Показатель", "Значение"], [("Вариант", "4"), ("Кодирование", "двоичное"), ("Простая задача", "P04"), ("Повышенная задача", "набор 7"), ("Контрольный вопрос", "4")])
    add_heading(doc, "1.2 Математическая постановка", level=2, page_break=False)
    add_body(doc, "Пусть x_i принимает значение 0 или 1 и показывает, выбран ли предмет i. Тогда требуется максимизировать стоимость F(x) = sum(p_i x_i) при ограничении sum(w_i x_i) <= C. Для P04 вместимость C равна 50, а эталонное решение имеет хромосому (1, 0, 0, 1, 0, 0, 0), суммарный вес 50 и стоимость 107.")
    add_body(doc, "Для набора 7 повышенной сложности используется 50 предметов и вместимость 12828. Точный эталонный результат, рассчитанный динамическим программированием, равен 16763 при весе 12820.")
    add_table(doc, "Таблица 2 - Базовые параметры генетического алгоритма", ["Параметр", "P04", "Набор 7"], [("Размер популяции", "60", "100"), ("Число поколений", "250", "400"), ("Вероятность кроссовера", "0,85", "0,85"), ("Вероятность мутации", "1/n", "0,02"), ("Метод кроссовера", "двухточечный", "двухточечный"), ("Элитизм", "2", "2"), ("Число seed", "5", "5")])

    add_heading(doc, "2 ТЕОРЕТИЧЕСКИЕ СВЕДЕНИЯ", level=1)
    add_heading(doc, "2.1 Двоичное кодирование", level=2, page_break=False)
    add_body(doc, "При двоичном кодировании хромосома имеет длину n и состоит из нулей и единиц. Позиция гена соответствует номеру предмета. Такое представление непосредственно задает подмножество предметов, однако стандартные операции кроссовера и мутации могут привести к превышению вместимости рюкзака.")
    add_heading(doc, "2.2 Генетические операторы и восстановление", level=2, page_break=False)
    add_body(doc, "В программе используется турнирный отбор, двухточечный кроссовер, побитовая мутация и элитизм. После формирования хромосомы выполняется жадное восстановление: из перегруженного решения удаляется выбранный предмет с минимальным отношением стоимости к весу. Процедура повторяется до получения допустимого решения.")
    add_heading(doc, "2.3 Критерии оценки", level=2, page_break=False)
    add_body(doc, "Качество особи оценивается суммарной стоимостью восстановленного решения. В экспериментах дополнительно измеряются суммарный вес, ошибка относительно точного оптимума, процент точности и время выполнения. Фиксация seed позволяет сравнивать варианты параметров на воспроизводимых запусках.")

    add_heading(doc, "3 РЕАЛИЗАЦИЯ И РЕЗУЛЬТАТЫ", level=1)
    add_heading(doc, "3.1 Структура программы", level=2, page_break=False)
    add_body(doc, "Файл data.py содержит проверяемые исходные данные P04 и набора 7. В knapsack_ga.py реализованы оценка и восстановление особей, точный динамический эталон, турнирный отбор, двухточечный кроссовер, мутация, элитизм и основной цикл генетического алгоритма. Сценарий run_experiments.py выполняет базовые серии и параметрическое исследование, сохраняет CSV/JSON и строит графики.")
    add_heading(doc, "3.2 Результаты базовых запусков", level=2, page_break=False)
    rows = []
    for name, item in (("P04", p04), ("Набор 7", set7)):
        rows.append((name, item["exact"]["profit"], item["exact"]["weight"], f"{item['baseline']['best_profit']:.0f}", f"{item['baseline']['mean_profit']:.1f}", f"{item['baseline']['mean_accuracy_percent']:.2f}", f"{item['baseline']['mean_time_seconds']:.3f}"))
    add_table(doc, "Таблица 3 - Сравнение с точным оптимумом", ["Набор", "Точный profit", "Вес", "Лучший ГА", "Средний ГА", "Точность, %", "Среднее время, с"], rows, font_size=9)
    add_body(doc, f"Для P04 генетический алгоритм во всех пяти запусках получил оптимальную стоимость 107 и вес 50. Для набора 7 лучшая серия достигла точного результата 16763 при среднем значении {set7['baseline']['mean_profit']:.1f}, что соответствует средней точности {set7['baseline']['mean_accuracy_percent']:.2f}%. Все сохраненные решения удовлетворяют ограничению вместимости.")
    add_image(doc, "convergence_p04.png", "Рисунок 1 - Сходимость генетического алгоритма для P04")
    add_image(doc, "convergence_набор_7.png", "Рисунок 2 - Сходимость генетического алгоритма для набора 7")

    add_heading(doc, "4 ИССЛЕДОВАНИЕ ПАРАМЕТРОВ", level=1)
    add_heading(doc, "4.1 Влияние мощности популяции", level=2, page_break=False)
    add_body(doc, "При увеличении популяции растет разнообразие хромосом и количество вычислений в каждом поколении. Для набора 7 сравнивались 20, 60 и 100 особей. Увеличение популяции не гарантирует монотонного роста точности, но снижает риск преждевременной потери перспективных комбинаций.")
    add_image(doc, "parameter_population.png", "Рисунок 3 - Влияние мощности популяции на среднюю стоимость")
    add_heading(doc, "4.2 Влияние вероятности кроссовера", level=2, page_break=False)
    add_body(doc, "Вероятность кроссовера определяет долю потомков, полученных путем обмена фрагментами родительских хромосом. При малой вероятности сильнее сохраняются родительские решения, а при значении, близком к единице, пространство поиска исследуется интенсивнее. В работе рассмотрены значения 0,60, 0,85 и 1,00.")
    add_image(doc, "parameter_crossover.png", "Рисунок 4 - Влияние вероятности кроссовера")
    add_heading(doc, "4.3 Влияние вероятности мутации", level=2, page_break=False)
    add_body(doc, "Мутация поддерживает разнообразие популяции и позволяет выходить из локальных областей поиска. Слишком малая вероятность ограничивает разнообразие, а слишком большая разрушает найденные комбинации. В работе рассмотрены значения 0,005, 0,02 и 0,05.")
    add_image(doc, "parameter_mutation.png", "Рисунок 5 - Влияние вероятности мутации")
    table_rows = []
    for label in ("population", "crossover", "mutation"):
        values = sorted({row["label"].split("=", 1)[1] for row in study if row["problem"] == "Набор 7" and row["label"].startswith(label + "=")}, key=float)
        for value in values:
            subset = [row for row in study if row["problem"] == "Набор 7" and row["label"] == f"{label}={value}"]
            table_rows.append((label, value, f"{mean(subset, 'profit'):.1f}", f"{mean(subset, 'accuracy_percent'):.2f}", f"{mean(subset, 'elapsed_seconds'):.3f}"))
    add_table(doc, "Таблица 4 - Результаты параметрического исследования набора 7", ["Фактор", "Значение", "Средняя стоимость", "Точность, %", "Время, с"], table_rows, font_size=9)

    add_heading(doc, "5 ВЫВОДЫ", level=1)
    add_body(doc, "В ходе лабораторной работы реализован генетический алгоритм с двоичным кодированием для задачи 0/1-рюкзака. Для устранения недопустимых решений применено жадное восстановление по отношению стоимости к весу.")
    add_body(doc, "На тесте P04 алгоритм во всех базовых запусках нашел эталонное решение со стоимостью 107 и весом 50. На наборе 7 повышенной сложности лучшая серия достигла точного динамического оптимума 16763 при допустимом весе 12820. Это подтверждает работоспособность двоичного кодирования и оператора восстановления.")
    add_body(doc, "Параметрическое исследование показало, что размер популяции и вероятности генетических операторов задают компромисс между разнообразием поиска, точностью и временем. Результаты экспериментов сохранены в CSV и JSON, а исходный код приведен в приложениях.")

    add_heading(doc, "6 ОТВЕТ НА КОНТРОЛЬНЫЙ ВОПРОС", level=1)
    add_body(doc, "Алгоритм восстановления преобразует недопустимую особь в допустимую, если после скрещивания или мутации общий вес выбранных предметов превышает вместимость рюкзака. Для этого из текущего решения последовательно удаляются выбранные предметы по заданному правилу. В данной работе удаляется предмет с минимальным отношением стоимости к весу. После устранения перегрузки решение снова проверяется, и полученная допустимая хромосома используется для оценки и дальнейшей эволюции.")

    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    refs = [
        "1. Лабораторная работа № 3. Решение задач комбинаторной оптимизации с помощью генетических алгоритмов на примере задачи укладки рюкзака: методические указания.",
        "2. Вирсански Э. Генетические алгоритмы на Python / пер. с англ. А. А. Слинкина. - М.: ДМК Пресс, 2020. - 286 с.",
        "3. Скобцов Ю. А. Вычислительный интеллект. - СПб.: ГУАП, 2022. - 138 с.",
        "4. GUAP. Нормативная документация для учебного процесса. - URL: https://guap.ru/c/regdocs/docs/uch (дата обращения: 27.09.2026).",
        "5. Burkardt J. KNAPSACK_01: Data for the 01 Knapsack Problem. - URL: https://people.sc.fsu.edu/~jburkardt/datasets/knapsack_01/ (дата обращения: 27.09.2026).",
    ]
    for ref in refs:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
        run = paragraph.add_run(ref)
        set_run_font(run, size=14)

    add_code_listing(doc, "knapsack_ga.py", "Приложение А. Листинг алгоритма")
    add_code_listing(doc, "data.py", "Приложение Б. Исходные данные")
    add_code_listing(doc, "run_experiments.py", "Приложение В. Сценарий экспериментов")
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
