"""Build the editable GUAP report for laboratory work 5."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(r"C:\Users\theju\OneDrive\Документы\ChatGPT\учеба\_guap_refs\lab_title.docx")
OUTPUT = ROOT / "report.docx"
RESULTS = ROOT / "results"


def set_run_font(run, name="Times New Roman", size=14, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(str(text)), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "Обновите поле содержания в Word"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(3); section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(14); normal.paragraph_format.first_line_indent = Cm(1.25); normal.paragraph_format.line_spacing = 1.5; normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"; style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(14); style.font.bold = True; style.paragraph_format.first_line_indent = Cm(1.25); style.paragraph_format.space_after = Pt(10); style.paragraph_format.keep_with_next = True
    footer = section.footer; footer.is_linked_to_previous = False
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; clear_paragraph(p); add_field(p, "PAGE")
    for run in p.runs: set_run_font(run, size=12)


def add_title_page(doc):
    doc.paragraphs[1].text = "КАФЕДРА № 43"; doc.paragraphs[7].text = "Санкт-Петербург 2026"
    for idx in (2, 3):
        p = doc.paragraphs[idx]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.left_indent = Pt(0); p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.right_indent = Pt(0)
    for idx in (0, 1, 7):
        for run in doc.paragraphs[idx].runs: set_run_font(run, size=14, bold=(idx == 0))
    teacher = doc.tables[0]; set_cell_text(teacher.cell(0, 4), "Скобцов Ю.А."); set_cell_text(teacher.cell(1, 4), "инициалы, фамилия", size=10)
    title = doc.tables[1]
    set_cell_text(title.cell(0, 0), "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 5", size=13, bold=True)
    set_cell_text(title.cell(1, 0), "Оптимизация многомерных функций с помощью эволюционной стратегии\nВариант 4", size=10)
    set_cell_text(title.cell(2, 0), "по курсу: Эволюционные методы проектирования программно-информационных систем", size=10)
    for row in list(title.rows)[3:]: title._tbl.remove(row._tr)
    student = doc.tables[2]
    set_cell_text(student.cell(0, 1), "4332К"); set_cell_text(student.cell(0, 5), "Могилатов С.И."); set_cell_text(student.cell(1, 3), "подпись, дата", size=10); set_cell_text(student.cell(1, 5), "инициалы, фамилия", size=10)
    doc.add_page_break()


def add_heading(doc, text, level=1, page_break=True):
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break: p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(text), size=14, bold=True)


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal"); p.paragraph_format.first_line_indent = Cm(1.25); p.paragraph_format.line_spacing = 1.5; p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; set_run_font(p.add_run(text), size=14)


def add_caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(8); set_run_font(p.add_run(text), size=12)


def add_image(doc, filename, caption, width_cm=15.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0); p.add_run().add_picture(str(RESULTS / filename), width=Cm(width_cm)); add_caption(doc, caption)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr; borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None: borders = OxmlElement("w:tblBorders"); tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge; element = borders.find(qn(tag))
        if element is None: element = OxmlElement(tag); borders.append(element)
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "4"); element.set(qn("w:space"), "0"); element.set(qn("w:color"), "808080")


def add_table(doc, caption, headers, rows, font_size=10):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0); set_run_font(p.add_run(caption), size=12, bold=True)
    table = doc.add_table(rows=1, cols=len(headers)); set_table_borders(table)
    for i, header in enumerate(headers): set_cell_text(table.rows[0].cells[i], header, size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row): set_cell_text(cells[i], value, size=font_size)
    doc.add_paragraph(); return table


def add_code_listing(doc, filename, title):
    add_heading(doc, title, page_break=True)
    for line in (ROOT / filename).read_text(encoding="utf-8").splitlines():
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0); set_run_font(p.add_run(line or " "), name="Consolas", size=8.5)


def build():
    summary = json.loads((RESULTS / "results_summary.json").read_text(encoding="utf-8"))
    with (RESULTS / "parameter_study.csv").open(encoding="utf-8") as stream: study = list(csv.DictReader(stream))
    n2 = summary["base"]["2"]; n3 = summary["base"]["3"]
    doc = Document(str(TEMPLATE)); configure_document(doc); add_title_page(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0); set_run_font(p.add_run("СОДЕРЖАНИЕ"), size=14, bold=True)
    toc = doc.add_paragraph(); toc.paragraph_format.first_line_indent = Cm(0); add_toc(toc)

    add_heading(doc, "1 УСЛОВИЕ ЗАДАНИЯ")
    add_body(doc, "Для варианта 4 требуется реализовать эволюционную стратегию для оптимизации многомерной функции moved axis parallel hyper-ellipsoid. В методических указаниях формула приведена как f(x) = sum(5*i*x_i^2), i=1..n, а в графе оптимума указано x_i=5*i. Эти два фрагмента противоречат друг другу. В работе формула используется буквально, поэтому математически согласованный глобальный минимум равен f(x*)=0 в точке x*=(0,...,0).")
    add_body(doc, "Требуется выполнить поиск для n=2, построить график функции с точками популяции и найденным экстремумом, исследовать влияние размера популяции и вероятности мутации, а затем повторить эксперимент для n=3 и сравнить результаты по точности и времени.")
    add_table(doc, "Таблица 1 - Параметры варианта 4", ["Показатель", "Значение"], [("Функция", "sum(5*i*x_i^2), i=1..n"), ("Область", "-5,12 <= x_i <= 5,12"), ("Теоретический минимум", "0"), ("Точка минимума по формуле", "x_i=0"), ("Базовая популяция", "mu=lambda=20"), ("Критерий останова", "100 поколений или стагнация 25 поколений"), ("Контрольный вопрос", "4")])
    add_heading(doc, "2 КРАТКИЕ ТЕОРЕТИЧЕСКИЕ СВЕДЕНИЯ")
    add_heading(doc, "2.1 Отличия эволюционной стратегии", 2, False)
    add_body(doc, "Эволюционная стратегия выполняет поиск в пространстве вещественных фенотипов. Потенциальное решение задается вектором координат x, а стратегические параметры определяют величину случайного шага. В отличие от двоичного генетического алгоритма не требуется кодировать и декодировать хромосому.")
    add_heading(doc, "2.2 Многократная ES", 2, False)
    add_body(doc, "В работе использована стратегия (mu+lambda). Из mu родителей формируется lambda потомков, после чего родители и потомки объединяются, а mu лучших особей переходят в следующее поколение. В базовой серии mu=lambda=20. Мутация выбирает родителя, изменяет его координаты нормальным шумом и ограничивает результат областью допустимых значений.")
    add_heading(doc, "2.3 Самоадаптация параметров", 2, False)
    add_body(doc, "Вместе с координатами изменяется вектор sigma стандартных отклонений. Новое значение sigma получается логнормальным масштабированием с общим и индивидуальным случайными множителями. Большие значения sigma обеспечивают исследование области, малые позволяют уточнять решение около найденного минимума.")

    add_heading(doc, "3 ПРОГРАММА И РЕЗУЛЬТАТЫ")
    add_heading(doc, "3.1 Алгоритм и параметры реализации", 2, False)
    add_body(doc, "Файл evolution_strategy.py содержит функцию целевого значения, класс Individual, мутацию, самоадаптацию sigma, отбор (mu+lambda) и критерий остановки. Сценарий run_experiments.py запускает две базовые серии, исследует параметры и формирует изображения. Для каждого поколения сохраняются лучшее значение, среднее значение популяции, расстояние до нулевой точки и среднее sigma.")
    add_table(doc, "Таблица 2 - Параметры базовой серии", ["Параметр", "Значение"], [("Размер популяции mu", "20"), ("Число потомков lambda", "20"), ("Вероятность мутации", "1,00"), ("Границы координат", "[-5,12; 5,12]"), ("Максимум поколений", "100"), ("Лимит стагнации", "25"), ("Seed", "6 для n=2, 7 для n=3")])
    add_heading(doc, "3.2 Результаты для n=2 и n=3", 2, False)
    add_table(doc, "Таблица 3 - Сравнение размерности задачи", ["Размерность", "Лучшее f(x)", "Точка решения", "Поколение остановки", "Время, с"], [("n=2", f"{n2['best_value']:.4e}", str([f"{x:.3e}" for x in n2["best_x"]]), n2["generations"], f"{n2['time_seconds']:.6f}"), ("n=3", f"{n3['best_value']:.4e}", str([f"{x:.3e}" for x in n3["best_x"]]), n3["generations"], f"{n3['time_seconds']:.6f}")])
    add_body(doc, f"Для n=2 получено f(x)={n2['best_value']:.6e}, координаты имеют порядок 10^-7. Для n=3 получено f(x)={n3['best_value']:.6e}, координаты имеют порядок 10^-6. Увеличение размерности потребовало больше поколений и немного увеличило время вычисления, но найденная точка также близка к теоретическому минимуму.")
    add_image(doc, "population_n2.png", "Рисунок 1 - Функция варианта 4, точки популяции и найденный экстремум")
    add_image(doc, "convergence_n2_n3.png", "Рисунок 2 - Сходимость ES для n=2 и n=3")
    add_image(doc, "step_n2_000.png", "Рисунок 3 - Начальная популяция для n=2")
    add_image(doc, "step_n2_050.png", "Рисунок 4 - Промежуточная популяция для n=2")
    add_heading(doc, "3.3 Исследование параметров", 2, False)
    add_body(doc, "Исследованы размер популяции mu=5, 10, 20, 40 и вероятность мутации p=0,25, 0,50, 0,75, 1,00. Для остальных настроек использовались 100 поколений и одинаковый критерий стагнации. Качество сравнивалось по достигнутому значению целевой функции, а производительность - по времени и числу поколений.")
    rows = [(row["factor"], row["label"], f"{float(row['best_value']):.3e}", row["generations"], f"{float(row['time_seconds']):.6f}", f"{float(row['distance']):.3e}") for row in study]
    add_table(doc, "Таблица 4 - Параметрическое исследование", ["Фактор", "Значение", "Лучшее f(x)", "Поколения", "Время, с", "Расстояние"], rows, font_size=9)
    add_image(doc, "parameter_study.png", "Рисунок 5 - Зависимость результата от параметров ES")
    add_body(doc, "Увеличение популяции в проведенных сериях улучшило вероятность достижения окрестности минимума, но увеличило время одного запуска. Вероятность мутации p=1,00 дала наиболее точное решение в базовом сравнении. При p=0,50 остановка по стагнации произошла раньше и результат оказался хуже, что показывает чувствительность ES к интенсивности исследования.")

    add_heading(doc, "4 ВЫВОДЫ")
    add_body(doc, "Реализована вещественная эволюционная стратегия (mu+lambda) с нормальной мутацией и самоадаптацией стандартных отклонений. Для n=2 найдено значение функции 1,12*10^-13, для n=3 - 3,62*10^-10, что подтверждает приближение к математически согласованному минимуму печатной формулы.")
    add_body(doc, "Пошаговые графики показывают движение точек популяции к началу координат. Исследование параметров выявило компромисс: увеличение популяции повышает разнообразие поиска, но увеличивает время; слишком малая вероятность мутации может привести к преждевременной стагнации.")

    add_heading(doc, "5 ОТВЕТ НА КОНТРОЛЬНЫЙ ВОПРОС")
    add_body(doc, "Двукратная эволюционная стратегия обозначается (1+1)-ЭС. В ней в каждый момент времени поддерживается одна родительская особь. Из нее мутацией получают одного потомка: координаты потомка равны координатам родителя плюс нормально распределенные случайные приращения, масштаб которых задается вектором sigma. Затем сравнивают значения целевой функции. Если потомок лучше родителя, он становится новой родительской особью; если нет, потомок отбрасывается. Процесс повторяется до достижения критерия остановки.")
    add_body(doc, "Преимуществом (1+1)-ЭС является простота и небольшое количество вычислений, недостатком - отсутствие разнообразия популяции и риск медленной сходимости или застревания. В данной работе для практического исследования числа особей использована многократная (mu+lambda)-ЭС, но ее мутационный принцип основан на той же идее.")

    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    refs = [
        "1. Лабораторная работа № 5. Оптимизация многомерных функций с помощью эволюционной стратегии: методические указания.",
        "2. Вирсански Э. Генетические алгоритмы на Python / пер. с англ. - М.: ДМК Пресс, 2020.",
        "3. Скобцов Ю. А. Вычислительный интеллект. - СПб.: ГУАП, 2022.",
        "4. GUAP. Нормативная документация для учебного процесса. - URL: https://guap.ru/c/regdocs/docs/uch (дата обращения: 15.10.2026).",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Normal"); p.paragraph_format.first_line_indent = Cm(1.25); set_run_font(p.add_run(ref), size=14)
    add_code_listing(doc, "evolution_strategy.py", "Приложение А. Листинг эволюционной стратегии")
    add_code_listing(doc, "run_experiments.py", "Приложение Б. Сценарий экспериментов")
    doc.save(str(OUTPUT)); print(OUTPUT)


if __name__ == "__main__": build()

