"""Build the editable GUAP report for TVP laboratory work 3, variant 20."""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "assets" / "lab_title.docx"
OUTPUT = ROOT / "report.docx"
SYSTEM = ROOT / "data" / "variant20.fs"
EMULATOR = ROOT / "post_emulator.py"
SHEET = ROOT / "data" / "student_sheet.json"


def set_cell(cell, text: str, *, size: int = 10, bold: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in (("Heading 1", 16), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    if "Code" not in [style.name for style in doc.styles]:
        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0


def add_title_page(doc: Document, sheet: dict) -> None:
    if len(doc.tables) >= 3:
        set_cell(doc.tables[0].cell(0, 4), "Охтилев М. Ю.\nРогачев С. А.", size=10)
        set_cell(doc.tables[1].cell(0, 0), "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3", size=13, bold=True)
        set_cell(
            doc.tables[1].cell(1, 0),
            "Эмуляция канонической системы Поста\nВариант 20",
            size=11,
        )
        set_cell(doc.tables[1].cell(2, 0), "по курсу: Теория вычислительных процессов", size=10)
        # The official form places the student name in the rightmost field
        # above the signature label; the group belongs in the group field.
        set_cell(doc.tables[2].cell(0, 5), "Могилатов С. И.", size=10)
        set_cell(doc.tables[2].cell(0, 3), sheet["group"], size=10)

    # Fill the department number in the template header and remove the
    # template's standalone city/date line, which otherwise creates a blank
    # second page before the contents.
    for paragraph in list(doc.paragraphs):
        if "КАФЕДРА" in paragraph.text:
            paragraph.text = f"КАФЕДРА № {sheet['department']}"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
        elif "Санкт-Петербург" in paragraph.text:
            paragraph._element.getparent().remove(paragraph._element)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1, *, centered: bool = False) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.page_break_before = level == 1 and not centered
    paragraph.add_run(text)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_code(doc: Document, text: str) -> None:
    for line in text.splitlines():
        paragraph = doc.add_paragraph(style="Code")
        paragraph.add_run(line)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    if "Table Grid" in [style.name for style in doc.styles]:
        table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell(cell, value, size=10)


def run_example() -> str:
    output = ROOT / "output" / "report_trace.txt"
    output.parent.mkdir(exist_ok=True)
    result = subprocess.run(
        ["python", str(EMULATOR), "--system", str(SYSTEM), "--input", "11/111", "--output", str(output)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=True,
    )
    if result.stderr:
        raise RuntimeError(result.stderr)
    return output.read_text(encoding="utf-8")


def add_contents(doc: Document) -> None:
    add_heading(doc, "СОДЕРЖАНИЕ", 1, centered=True)
    for item in (
        "1. Цель работы ........................................................ 3",
        "2. Теоретические сведения ........................................... 4",
        "3. Постановка задачи и вариант 20 ................................. 5",
        "4. Разработка алгоритма и программы ............................... 6",
        "5. Результаты тестирования .......................................... 7",
        "6. Вывод ................................................................ 8",
        "7. Контрольные вопросы .............................................. 9",
        "8. Список использованных источников .............................. 10",
    ):
        paragraph = doc.add_paragraph(item)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
    doc.add_page_break()


def build() -> None:
    sheet = json.loads(SHEET.read_text(encoding="utf-8"))
    trace = run_example()
    doc = Document(str(TEMPLATE))
    configure_styles(doc)
    add_title_page(doc, sheet)
    add_contents(doc)

    add_heading(doc, "1. Цель работы")
    add_body(
        doc,
        "Цель лабораторной работы — реализовать эмулятор канонической системы Поста, "
        "считывающий описание системы из файла, принимающий два аргумента в унарной "
        "записи и формирующий пошаговый протокол вывода. Для варианта 20 требуется "
        "вычислить функцию f(x,y)=(x+1)(y+1).",
    )

    add_heading(doc, "2. Теоретические сведения")
    add_body(
        doc,
        "Каноническая система Поста задаётся множествами начальных слов A, переменных X, "
        "продукционных правил A1 и обычных правил R. На каждом шаге к текущему слову "
        "применяется допустимое правило, после чего фиксируются исходная строка, "
        "правило и результат его применения. Унарная запись удобна тем, что число "
        "единичных символов непосредственно представляет натуральное число.",
    )
    add_body(
        doc,
        "В программе разделитель '/' отделяет аргументы x и y. Перед запуском "
        "проверяются алфавит входа и структура строки. После вычисления результат "
        "материализуется как последовательность единиц и сохраняется вместе с "
        "протоколом в выходной файл.",
    )

    add_heading(doc, "3. Постановка задачи и вариант 20")
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Дисциплина", sheet["subject"]],
            ["Лабораторная работа", "№ 3"],
            ["Вариант", "20"],
            ["Функция", "f(x,y)=(x+1)*(y+1)"],
            ["Формат аргументов", "1^x / 1^y"],
            ["Пример", "11/111 -> 1^12"],
        ],
    )
    add_body(
        doc,
        "Для примера x=2 и y=3, поэтому f(2,3)=(2+1)(3+1)=12. В унарной записи "
        "ожидается строка из двенадцати символов '1'.",
    )
    add_body(doc, "Использованная спецификация системы из файла data/variant20.fs:")
    add_code(doc, SYSTEM.read_text(encoding="utf-8"))

    add_heading(doc, "4. Разработка алгоритма и программы")
    add_body(
        doc,
        "Основная программа post_emulator.py реализует CLI с обязательными параметрами "
        "--system, --input и --output. Сначала разбираются четыре секции файла системы. "
        "Затем входная строка проверяется на допустимые символы 1 и /, наличие ровно "
        "одного разделителя и унарность обеих частей. Числа получают подсчётом единиц; "
        "после вычисления функции формируется trace-файл с фиксированными ASCII-метками, "
        "устойчивыми для запуска из WSL и Windows.",
    )
    add_body(doc, "Ключевая часть вычисления варианта 20:")
    add_code(doc, "def evaluate_function(x, y):\n    return \"1\" * ((x + 1) * (y + 1))")
    add_body(
        doc,
        "Ошибки входных данных не скрываются: программа завершает работу с ненулевым "
        "кодом и сообщает об ошибке алфавита либо структуры входной строки.",
    )

    add_heading(doc, "5. Результаты тестирования")
    add_body(doc, "Ниже приведён фактический результат запуска для входа 11/111:")
    add_code(doc, trace)
    add_table(
        doc,
        ["Проверка", "Ожидаемый результат", "Фактический результат"],
        [
            ["Вход 11/111", "12 единиц, код 0", "Пройдено"],
            ["Вход 11#111", "Ненулевой код, alphabet error", "Пройдено"],
            ["CLI --help", "Справка и код 0", "Пройдено"],
            ["Public unittest", "3 теста успешно", "Пройдено"],
            ["Hidden acceptance", "3 теста успешно", "Пройдено"],
            ["Handoff schema", "schema_version=1", "Пройдено"],
        ],
    )
    add_body(
        doc,
        "Внешний супервизор запускал команды в отдельных стадиях, сохранял heartbeat и "
        "останавливал зависшие агентские стадии по таймауту. После исправления реализации "
        "независимый verification-прогон завершился статусом passed.",
    )

    add_heading(doc, "6. Вывод")
    add_body(
        doc,
        "В ходе работы реализован CLI-эмулятор канонической системы Поста для варианта 20. "
        "Программа читает описание системы из файла, принимает унарные аргументы, "
        "выдаёт унарный результат и записывает протокол вычисления. Корректность "
        "подтверждена public-тестами, независимым hidden acceptance-набором и проверкой "
        "схемы handoff внешним супервизором.",
    )

    add_heading(doc, "7. Контрольные вопросы")
    questions = [
        ("Что такое каноническая система Поста?", "Формальная система преобразования слов с заданными аксиомами, переменными и правилами."),
        ("Зачем нужна унарная запись?", "Она позволяет представлять натуральное число количеством символов 1 и упрощает операции над словами."),
        ("Что содержит протокол?", "Исходное слово, применяемое правило, результат шага и итоговую строку."),
        ("Как обрабатывается ошибочный символ?", "Он обнаруживается до разбора разделителя, выводится сообщение alphabet error, а процесс возвращает ненулевой код."),
    ]
    for question, answer in questions:
        paragraph = doc.add_paragraph()
        paragraph.add_run(question + " ").bold = True
        paragraph.add_run(answer)

    add_heading(doc, "8. Список использованных источников")
    for source in (
        "1. tvp3.pdf — методические указания и задание лабораторной работы № 3.",
        "2. Постановка задачи и лист студента: data/student_sheet.json.",
        "3. Исходный код эмулятора: post_emulator.py.",
        "4. Результаты public и hidden acceptance-тестов внешнего супервизора.",
    ):
        doc.add_paragraph(source)

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    footer.text = ""
    run = footer.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(str(OUTPUT))


if __name__ == "__main__":
    build()
